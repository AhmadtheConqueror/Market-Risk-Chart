"""Synthetic test assessments only; no proprietary workbook is bundled."""
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from openpyxl import Workbook, load_workbook
import pytest
from sqlalchemy import select, delete, func
from app.models.historical import HistoricalMarketObservation as H
from app.services.historical_market import HISTORICAL_MAPPING, ANALYTICAL_MAPPING, RISK_SCORING_KEYS
from app.services.official_market import load_snapshot_history
from app.services.ai_context import build_dashboard_ai_context


def workbook_bytes():
    w=Workbook();s=w.active;s.title="Market Risk";s.cell(21,6,"AssessDate")
    dates=[];day=date(2026,10,14)
    while len(dates)<199:
        if day.weekday()<5:dates.append(day)
        day-=timedelta(days=1)
    for col,(key,(code,label,unit,*_)) in enumerate(HISTORICAL_MAPPING.items(),7):
        for row,value in [(17,code),(18,label),(19,unit.split("/")[1]),(20,"USD"),(21,"c")]:s.cell(row,col,value)
        for i,day in enumerate(sorted(dates),22):
            s.cell(i,6,datetime.combine(day,datetime.min.time()))
            s.cell(i,col,float(i+col))
    output=BytesIO();w.save(output);return output.getvalue()


@pytest.fixture
def imported_sixteen(client, seed_test_data):
    payload=workbook_bytes()
    result=client.post("/api/market/import-excel",files={"workbook":("SYNTHETIC_16.xlsx",payload)})
    assert result.status_code==200,result.text
    assert result.json()["rows_valid"]==16*199
    return payload


def test_exact_analytical_codes_units_and_fixed_scoring_membership():
    assert {v[0] for v in ANALYTICAL_MAPPING.values()}=={
        "PCAIC00","PCNGC00","AAEJB00","AAXUO00","AAEIZ00","PCAID00","AAYWT00","AAWZD00","PJAAU00"}
    assert RISK_SCORING_KEYS=={"brent","naphtha","gasoil","gasoline","forcados","wti","jet"}
    assert len(HISTORICAL_MAPPING)==16
    assert all(v[2]==("USD/MT" if k in {"gasoil_ara","ulsd_nwe","jet_cif"} else "USD/BBL")
               and v[3] is None for k,v in ANALYTICAL_MAPPING.items())


def test_sixteen_latest_history_statistics_and_advisor_boundary(client, imported_sixteen, db_session):
    snapshot=client.get("/api/dashboard/snapshot").json()
    historical={s["instrument"]:s for s in client.get("/api/market/historical").json()["instruments"]}
    latest={o["instrumentId"]:o for o in client.get("/api/market/latest").json()["observations"]}
    assert len(snapshot["market_stats"])==len(historical)==len(latest)==16
    for s in snapshot["market_stats"]:
        key=s["instrument_id"];h=historical[key]
        assert (s["current_value"],s["latest_date"],s["unit"])==(h["latest"],h["latest_date"],h["unit"])
        assert h["observation_count"]==199 and s["window_count"]==h["count_90d"]>=60
        assert s["z_score"]==h["z_score"]
        assert s["included_in_risk_scoring"]==(key in RISK_SCORING_KEYS)
        assert s["monitored_for_analytics"]
        if key in {"gasoil_ara","ulsd_nwe","jet_cif"}:
            assert latest[key]["convertedValue"] is None and h["spread"] is None and h["factor"] is None
    assert {s["instrument_id"] for s in snapshot["product_spreads"]}=={"naphtha","gasoil","gasoline","jet"}
    context=build_dashboard_ai_context(db_session)
    assert {s["instrument_id"] for s in context["market_stats"]}==RISK_SCORING_KEYS
    assert all(s["instrument_id"] in RISK_SCORING_KEYS for s in context["outliers"])
    assert sum(map(len,load_snapshot_history(db_session).values()))<16*199


def test_sixteen_reimport_is_idempotent_and_retains_history(client, imported_sixteen, db_session):
    before=db_session.scalar(select(func.count()).select_from(H))
    result=client.post("/api/market/import-excel",files={"workbook":("SYNTHETIC_16.xlsx",imported_sixteen)}).json()
    assert result["historical_rows_inserted"]==0
    assert result["historical_rows_updated_or_unchanged"]==16*199
    assert db_session.scalar(select(func.count()).select_from(H))==before


def test_invalid_or_absent_new_series_is_isolated(client, imported_sixteen, db_session):
    w=load_workbook(BytesIO(imported_sixteen));s=w["Market Risk"]
    for col,(key,spec) in enumerate(HISTORICAL_MAPPING.items(),7):
        if key=="bonny_light":s.cell(21,col,"u")
        if key=="bonga":s.cell(19,col,"INVALID")
        if key=="erha":s.cell(17,col).value=None
    output=BytesIO();w.save(output)
    result=client.post("/api/market/import-excel",files={"workbook":("SYNTHETIC_invalid.xlsx",output.getvalue())})
    assert result.status_code==200,result.text
    assert not {"bonny_light","bonga","erha"}&set(result.json()["mapped_historical_instruments"])
    assert db_session.scalar(select(func.count()).select_from(H))==16*199


def test_new_series_outlier_and_insufficient_history(client, imported_sixteen, db_session):
    rows=db_session.scalars(select(H).where(H.instrument_key=="erha").order_by(H.assessment_date)).all()
    rows[-1].native_value=100000;db_session.flush()
    outliers=client.get("/api/dashboard/snapshot").json()["outliers"]
    assert "erha" in {s["instrument_id"] for s in outliers}
    assert all(abs(outliers[i]["z_score"])>=abs(outliers[i+1]["z_score"]) for i in range(len(outliers)-1))
    db_session.execute(delete(H).where(H.instrument_key=="erha",H.id.not_in([r.id for r in rows[-59:]])))
    result=client.get("/api/dashboard/snapshot").json()
    stat=next(s for s in result["market_stats"] if s["instrument_id"]=="erha")
    assert stat["history_status"]=="insufficient_history" and stat["z_score"] is None
    assert "erha" not in {s["instrument_id"] for s in result["outliers"]}


def test_new_native_missing_values_are_omitted_and_zero_preserved(client):
    w=load_workbook(BytesIO(workbook_bytes()));s=w["Market Risk"]
    for col,(key,spec) in enumerate(HISTORICAL_MAPPING.items(),7):
        if key=="bonny_light":s.cell(220,col).value=None
        if key=="bonga":s.cell(220,col,0)
    output=BytesIO();w.save(output)
    response=client.post("/api/market/import-excel",files={"workbook":("SYNTHETIC_null.xlsx",output.getvalue())})
    assert response.status_code==200
    series={r["instrument"]:r for r in client.get("/api/market/historical").json()["instruments"]}
    assert series["bonny_light"]["observation_count"]==198
    assert series["bonny_light"]["latest_date"]=="2026-10-13"
    assert series["bonga"]["latest"]==0 and series["bonga"]["latest_date"]=="2026-10-14"
