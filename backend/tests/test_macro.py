from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal
import io
from unittest.mock import AsyncMock, MagicMock, patch
import zipfile

from fastapi.testclient import TestClient
import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.db.base import Base
from app.db.session import get_db
from app.main import app
from app.models.macro import MacroIndicator
from app.providers.macro.base import MacroDataProvider, MacroObservation
from app.providers.macro.nbs import NBSMacroProvider
from app.providers.macro.nuprc import NUPRCProductionProvider
from app.providers.macro.spglobal_pmi import SPGlobalPMIProvider
from app.services.macro_ingestion import (
    CANONICAL_MACRO_CONFIG,
    determine_freshness,
    get_latest_macro_indicators,
    get_macro_history,
    run_macro_refresh,
    upsert_macro_observation,
    validate_macro_observation,
)


@pytest.fixture
def in_memory_db():
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()
        Base.metadata.drop_all(bind=engine)


@pytest.fixture
def client(in_memory_db):
    def override_get_db():
        try:
            yield in_memory_db
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def create_mock_cpi_workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Table1"

    # Row 1-3: Headers
    ws.cell(1, 1, "Table 1 Composite Consumer Price Index (Base: 2024 = 100)")
    ws.cell(2, 2, "")
    ws.cell(2, 3, "All Items Index")
    ws.cell(2, 13, "Food")
    ws.cell(2, 18, "All Items less Farm Produce and Energy\n(CORE)")

    ws.cell(3, 2, "Month")
    ws.cell(3, 6, "Year-on change (%)")
    ws.cell(3, 16, "Year-on change (%)")
    ws.cell(3, 21, "Year-on change (%)")

    # Row 4: Weights
    ws.cell(4, 1, "Weights")

    # Row 5: August 2026 data
    ws.cell(5, 1, 2026)
    ws.cell(5, 2, "August")
    ws.cell(5, 6, 15.3908)
    ws.cell(5, 16, 19.5672)
    ws.cell(5, 21, 13.2915)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def create_mock_gdp_workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "real gdp growth rate %"

    ws.cell(3, 1, "Real Growth Rate (%)")
    ws.cell(6, 2, 2026)
    ws.cell(6, 3, "Q2")
    ws.cell(7, 2, "Q1")
    ws.cell(7, 3, "Q2")

    # Row 10: Agriculture
    ws.cell(10, 1, "AGRICULTURE")
    ws.cell(10, 2, 1.25)
    ws.cell(10, 3, 2.10)

    # Row 20: Total Real GDP
    ws.cell(20, 1, "GDP at 2019 constant price")
    ws.cell(20, 2, 3.8850)
    ws.cell(20, 3, 4.4337)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# 1. NBS Dynamic Discovery Tests
def test_nbs_cpi_catalog_discovery():
    mock_html = """
    <html><body>
        <div class="colx resource alternate">
            <span class="resource-info" id="1435">August 2026 CPI Report</span>
            <a href="https://microdata.nigerianstat.gov.ng/index.php/catalog/154/download/1435">Download</a>
            <table><tr><td class="caption">Date</td><td>2026-09-15</td></tr></table>
        </div>
        <div class="colx resource">
            <span class="resource-info" id="1432">July 2026 CPI Report</span>
            <a href="https://microdata.nigerianstat.gov.ng/index.php/catalog/154/download/1432">Download</a>
            <table><tr><td class="caption">Date</td><td>2026-08-17</td></tr></table>
        </div>
    </body></html>
    """
    provider = NBSMacroProvider()
    report = provider.discover_latest_report(mock_html, catalog_id=154)
    assert report["resource_id"] == "1435"
    assert "1435" in report["download_url"]
    assert report["title"] == "August 2026 CPI Report"
    assert report["published_date"] == datetime(2026, 9, 15, tzinfo=timezone.utc)


def test_nbs_gdp_catalog_discovery():
    mock_html = """
    <html><body>
        <div class="colx resource alternate">
            <span class="resource-info" id="1433">Nigerian Gross Domestic Product Report Q2 2026</span>
            <a href="https://microdata.nigerianstat.gov.ng/index.php/catalog/147/download/1433">Download</a>
            <table><tr><td class="caption">Date</td><td>2026-08-31</td></tr></table>
        </div>
        <div class="colx resource">
            <span class="resource-info" id="1408">Nigerian Gross Domestic Product Report Q1 2026</span>
            <a href="https://microdata.nigerianstat.gov.ng/index.php/catalog/147/download/1408">Download</a>
            <table><tr><td class="caption">Date</td><td>2026-05-25</td></tr></table>
        </div>
    </body></html>
    """
    provider = NBSMacroProvider()
    report = provider.discover_latest_report(mock_html, catalog_id=147)
    assert report["resource_id"] == "1433"
    assert "1433" in report["download_url"]
    assert report["published_date"] == datetime(2026, 8, 31, tzinfo=timezone.utc)


# 2. NBS CPI Extraction Tests
def test_nbs_cpi_workbook_extraction():
    xlsx_bytes = create_mock_cpi_workbook_bytes()
    provider = NBSMacroProvider()
    obs_list = provider.parse_cpi_workbook(xlsx_bytes, source_url="https://test.url/cpi.zip")

    assert len(obs_list) == 3
    by_key = {o.indicator_key: o for o in obs_list}

    h = by_key["headline_inflation"]
    assert h.value == Decimal("15.3908")
    assert h.unit == "%"
    assert h.reporting_period == "Aug 2026"
    assert h.source == "National Bureau of Statistics Nigeria"

    f = by_key["food_inflation"]
    assert f.value == Decimal("19.5672")
    assert f.unit == "%"
    assert f.reporting_period == "Aug 2026"

    c = by_key["core_inflation"]
    assert c.value == Decimal("13.2915")
    assert c.unit == "%"
    assert c.reporting_period == "Aug 2026"


# 3. NBS GDP Extraction Tests
def test_nbs_gdp_workbook_extraction():
    xlsx_bytes = create_mock_gdp_workbook_bytes()
    provider = NBSMacroProvider()
    obs_list = provider.parse_gdp_workbook(xlsx_bytes, source_url="https://test.url/gdp.zip")

    assert len(obs_list) == 1
    gdp = obs_list[0]
    assert gdp.indicator_key == "real_gdp_growth"
    assert gdp.value == Decimal("4.4337")
    assert gdp.unit == "%"
    assert gdp.reporting_period == "Q2 2026"
    assert gdp.source == "National Bureau of Statistics Nigeria"
    assert gdp.metadata_json["definition"] == "total_real_gdp_growth_yoy"


# 4. Fail-closed Tests
def test_nbs_malformed_cpi_workbook_fails_closed():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Table1"
    ws.cell(1, 1, "Empty or unexpected layout")
    buf = io.BytesIO()
    wb.save(buf)

    provider = NBSMacroProvider()
    with pytest.raises(ValueError, match="NBS CPI extraction failed: expected labels not found"):
        provider.parse_cpi_workbook(buf.getvalue(), source_url="test")


def test_nbs_malformed_gdp_workbook_fails_closed():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "random_sheet"
    buf = io.BytesIO()
    wb.save(buf)

    provider = NBSMacroProvider()
    with pytest.raises(ValueError, match="NBS GDP extraction failed"):
        provider.parse_gdp_workbook(buf.getvalue(), source_url="test")


# 5. S&P Global / Stanbic IBTC PMI Tests
def test_spglobal_pmi_extraction():
    provider = SPGlobalPMIProvider()

    # Mock pypdf reader extraction
    mock_page1 = MagicMock()
    mock_page1.extract_text.return_value = "1 September 2026\nThe Stanbic IBTC Bank Nigeria PMI\nAugust 2026"
    mock_page2 = MagicMock()
    mock_page2.extract_text.return_value = "Key findings\nOverview"
    mock_page3 = MagicMock()
    mock_page3.extract_text.return_value = "August 2026\nReturn to contents\n54.3\nThe Stanbic IBTC Bank Nigeria PMI"

    with patch("pypdf.PdfReader") as mock_reader_cls:
        mock_reader = MagicMock()
        mock_reader.pages = [mock_page1, mock_page2, mock_page3]
        mock_reader_cls.return_value = mock_reader

        obs = provider.parse_pmi_pdf(b"%PDF-1.4 mock", source_url="https://test.url/pmi.pdf")
        assert obs.indicator_key == "nigeria_pmi"
        assert obs.value == Decimal("54.3")
        assert obs.unit == "index"
        assert obs.reporting_period == "Aug 2026"
        assert obs.source == "Stanbic IBTC Bank / S&P Global"
        assert obs.published_at == datetime(2026, 9, 1, tzinfo=timezone.utc)


def test_spglobal_pmi_invalid_range_rejected():
    provider = SPGlobalPMIProvider()

    mock_page = MagicMock()
    mock_page.extract_text.return_value = "August 2026\nheadline PMI posted 154.3\n"

    with patch("pypdf.PdfReader") as mock_reader_cls:
        mock_reader = MagicMock()
        mock_reader.pages = [mock_page]
        mock_reader_cls.return_value = mock_reader

        with pytest.raises(ValueError, match="Sanity check failed for nigeria_pmi"):
            provider.parse_pmi_pdf(b"%PDF-1.4 mock", source_url="https://test.url/pmi.pdf")


def test_spglobal_pmi_malformed_pdf_fails_closed():
    provider = SPGlobalPMIProvider()

    with patch("pypdf.PdfReader", side_effect=Exception("not a pdf")):
        with pytest.raises(ValueError, match="PMI extraction failed: malformed PDF publication"):
            provider.parse_pmi_pdf(b"<!DOCTYPE html>", source_url="https://test.url/pmi.pdf")


def test_spglobal_pmi_discovery_skips_html_response():
    provider = SPGlobalPMIProvider()
    pdf_bytes = b"%PDF-1.4\n" + (b"x" * 10001)
    html_bytes = b"<!DOCTYPE html>" + (b"x" * 10001)

    class MockResponse:
        def __init__(self, content: bytes, content_type: str) -> None:
            self.status_code = 200
            self.content = content
            self.headers = {"content-type": content_type}

    class MockClient:
        def __init__(self) -> None:
            self.calls: list[str] = []

        def get(self, url: str) -> MockResponse:
            self.calls.append(url)
            if len(self.calls) == 1:
                return MockResponse(html_bytes, "text/html")
            return MockResponse(pdf_bytes, "application/pdf")

    client = MockClient()
    content, url = provider.discover_latest_pdf(client)  # type: ignore[arg-type]

    assert content == pdf_bytes
    assert url == client.calls[1]
    assert len(client.calls) == 2


# 6. NUPRC Strict Crude Oil Tests
def test_nuprc_strict_crude_extraction():
    provider = NUPRCProductionProvider()
    mock_article_html = """
    <html><body>
    <title>Production Rises by 0.4% in August as Nigeria Meets OPEC Quota | NUPRC</title>
    <p>Nigeria produced 1, 677,777 barrels of crude oil and condensate per day in the month of August 2026, representing an increase of 0.4% when compared to the month of July.</p>
    <p>In strict crude oil terms (excluding condensate), Nigeria produced 1,500,190 barrels per day as the country met its OPEC quota for the 4th consecutive month.</p>
    </body></html>
    """
    obs = provider.parse_article_content(mock_article_html, source_url="https://nuprc.gov.ng/test")

    assert obs.indicator_key == "crude_oil_production"
    assert obs.value == Decimal("1.50019")
    assert obs.unit == "mbpd"
    assert obs.reporting_period == "Aug 2026"
    assert obs.source == "Nigerian Upstream Petroleum Regulatory Commission"
    # Strict crude only definition
    assert obs.metadata_json["definition"] == "crude_only_excluding_condensate"
    assert obs.metadata_json["raw_value"] == 1500190
    # Combined is preserved in metadata, NOT main card
    assert obs.metadata_json["crude_plus_condensate_bpd"] == 1677777
    assert obs.metadata_json["crude_plus_condensate_mbpd"] == 1.677777


def test_nuprc_missing_label_fails_closed():
    provider = NUPRCProductionProvider()
    mock_html = "<p>Nigeria produced some barrels of oil.</p>"
    with pytest.raises(ValueError, match="NUPRC extraction failed: strict crude oil production"):
        provider.parse_article_content(mock_html, source_url="test")


# 7. Persistence & Idempotence Tests
def test_upsert_idempotence(in_memory_db):
    obs = MacroObservation(
        indicator_key="headline_inflation",
        display_name="Headline Inflation",
        value=Decimal("15.3908"),
        unit="%",
        reporting_period="Aug 2026",
        source="National Bureau of Statistics Nigeria",
        source_url="https://test.url",
    )

    # 1. First insert
    rec1, status1 = upsert_macro_observation(in_memory_db, obs)
    assert status1 == "stored"
    assert rec1.id is not None

    # 2. Second insert with identical key, period, source
    rec2, status2 = upsert_macro_observation(in_memory_db, obs)
    assert status2 == "unchanged"
    assert rec2.id == rec1.id

    # Verify no duplicates
    count = in_memory_db.query(MacroIndicator).count()
    assert count == 1


def test_official_revision_handling(in_memory_db):
    obs = MacroObservation(
        indicator_key="real_gdp_growth",
        display_name="Real GDP Growth",
        value=Decimal("4.4337"),
        unit="%",
        reporting_period="Q2 2026",
        source="National Bureau of Statistics Nigeria",
    )
    rec1, status1 = upsert_macro_observation(in_memory_db, obs)
    assert status1 == "stored"

    # Publisher releases revised figure for same period
    revised_obs = MacroObservation(
        indicator_key="real_gdp_growth",
        display_name="Real GDP Growth",
        value=Decimal("4.5100"),
        unit="%",
        reporting_period="Q2 2026",
        source="National Bureau of Statistics Nigeria",
    )
    rec2, status2 = upsert_macro_observation(in_memory_db, revised_obs)
    assert status2 == "updated"
    assert rec2.id == rec1.id
    assert float(rec2.value) == 4.51
    assert rec2.status == "revised"
    assert rec2.metadata_json["revised"] is True
    assert len(rec2.metadata_json["revision_history"]) == 1


# 8. Fault Isolation in Refresh Service
@pytest.mark.anyio
async def test_fault_isolation_in_refresh_service(in_memory_db):
    class GoodProvider(MacroDataProvider):
        @property
        def provider_name(self) -> str:
            return "GoodProvider"

        async def fetch_latest(self) -> list[MacroObservation]:
            return [
                MacroObservation(
                    indicator_key="headline_inflation",
                    display_name="Headline Inflation",
                    value=Decimal("15.39"),
                    unit="%",
                    reporting_period="Aug 2026",
                    source="National Bureau of Statistics Nigeria",
                )
            ]

    class FailingProvider(MacroDataProvider):
        @property
        def provider_name(self) -> str:
            return "FailingProvider"

        async def fetch_latest(self) -> list[MacroObservation]:
            raise ConnectionError("Official publisher site timed out")

    summary = await run_macro_refresh(in_memory_db, providers=[GoodProvider(), FailingProvider()])
    assert summary.stored == 1
    assert summary.failed == 1
    assert len(summary.errors) == 1
    assert summary.errors[0]["provider"] == "FailingProvider"


# 9. Freshness Calculation Tests
def test_freshness_determination():
    as_of = datetime(2026, 9, 24, tzinfo=timezone.utc)

    # August 2026 monthly data evaluated on 2026-09-24 is fresh (~27 days old)
    assert determine_freshness("headline_inflation", "Aug 2026", as_of=as_of) == "fresh"

    # May 2026 monthly data evaluated on 2026-09-24 is stale (~118 days old > 75 days)
    assert determine_freshness("headline_inflation", "May 2026", as_of=as_of) == "stale"

    # Q2 2026 GDP (ends June 28) evaluated on 2026-09-24 is fresh (~88 days old <= 180 days)
    assert determine_freshness("real_gdp_growth", "Q2 2026", as_of=as_of) == "fresh"

    # Q4 2025 GDP (ends Dec 28) evaluated on 2026-09-24 is stale (~270 days old > 180 days)
    assert determine_freshness("real_gdp_growth", "Q4 2025", as_of=as_of) == "stale"


# 10. Honest Null != Zero Tests
def test_null_is_not_zero(in_memory_db):
    results = get_latest_macro_indicators(in_memory_db)
    assert len(results) == 6
    for r in results:
        assert r.value is None
        assert r.freshness_status == "unavailable"
        assert r.reporting_period == "Unavailable"


# 11. API Endpoints Tests
def test_macro_api_endpoints(client, in_memory_db):
    # Seed 1 observation
    obs = MacroObservation(
        indicator_key="headline_inflation",
        display_name="Headline Inflation",
        value=Decimal("15.39"),
        unit="%",
        reporting_period="Aug 2026",
        source="National Bureau of Statistics Nigeria",
    )
    upsert_macro_observation(in_memory_db, obs)

    # Test GET /api/macro/latest
    resp = client.get("/api/macro/latest")
    assert resp.status_code == 200
    data = resp.json()
    assert len(data) == 6
    by_key = {d["indicator_key"]: d for d in data}
    assert by_key["headline_inflation"]["value"] == 15.39
    assert by_key["headline_inflation"]["reporting_period"] == "Aug 2026"
    assert by_key["food_inflation"]["value"] is None

    # Test GET /api/macro/history
    resp_hist = client.get("/api/macro/history?indicator=headline_inflation")
    assert resp_hist.status_code == 200
    hist_data = resp_hist.json()
    assert len(hist_data) == 1
    assert hist_data[0]["indicator_key"] == "headline_inflation"

    # Test GET /api/macro/history invalid indicator
    resp_bad = client.get("/api/macro/history?indicator=unknown_xyz")
    assert resp_bad.status_code == 400


def test_macro_refresh_endpoint_uses_provider_contract(client):
    class GoodProvider(MacroDataProvider):
        @property
        def provider_name(self) -> str:
            return "GoodProvider"

        async def fetch_latest(self) -> list[MacroObservation]:
            return [
                MacroObservation(
                    indicator_key="headline_inflation",
                    display_name="Headline Inflation",
                    value=Decimal("15.39"),
                    unit="%",
                    reporting_period="Aug 2026",
                    source="National Bureau of Statistics Nigeria",
                )
            ]

    with patch("app.services.macro_ingestion.get_macro_providers", return_value=[GoodProvider()]):
        resp = client.post("/api/macro/refresh")

    assert resp.status_code == 200
    data = resp.json()
    assert data["requested"] == 6
    assert data["stored"] == 1
    assert data["failed"] == 0
    assert data["indicators"][0]["indicator_key"] == "headline_inflation"


def test_dashboard_snapshot_includes_macro(client, in_memory_db):
    resp = client.get("/api/dashboard/snapshot")
    assert resp.status_code == 200
    data = resp.json()
    assert "macro_indicators" in data
    assert len(data["macro_indicators"]) == 6
    keys = [m["indicator_key"] for m in data["macro_indicators"]]
    assert set(keys) == set(CANONICAL_MACRO_CONFIG.keys())
