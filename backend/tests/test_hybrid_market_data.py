from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from io import BytesIO

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from app.models.market import MarketInstrument, MarketObservation
from app.services.ai_context import build_dashboard_ai_context
from app.services.calculations import calculate_instrument_stats
from app.services.excel_ingestion import ExcelValidationError, parse_market_workbook
from app.services.source_resolution import resolve_observations


EXCEL_COLUMNS = {
    "PAAAM00": (7, "Naphtha FOB Rdam Barge $/mt", "MT", "USD"),
    "AAVJI00": (8, "Gasoil 0.1%S FOB Med Cargo", "MT", "USD"),
    "PGABM00": (9, "Gasoline Prem Unleaded 10ppmS FOB AR Barge", "MT", "USD"),
    "PCABC00": (10, "Forcados FOB Nigeria", "BBL", "USD"),
    "PJAAV00": (11, "Jet FOB NWE Cargo", "MT", "USD"),
}


def workbook_bytes(include_sheet: bool = True, invalid_date: bool = False, days: int = 2) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Market Risk" if include_sheet else "Wrong Sheet"
    if not include_sheet:
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    for symbol, (column, label, unit, currency) in EXCEL_COLUMNS.items():
        worksheet.cell(17, column, symbol)
        worksheet.cell(18, column, label)
        worksheet.cell(19, column, unit)
        worksheet.cell(20, column, currency)
    worksheet.cell(21, 6, "AssessDate")
    for offset in range(days):
        row = 22 + offset
        worksheet.cell(row, 6, "not-a-date" if invalid_date and offset == 0 else datetime(2026, 9, 10) - timedelta(days=offset))
        for symbol, (column, _, _, _) in EXCEL_COLUMNS.items():
            value = 0.0 if (symbol == "PCABC00" and offset == 0) else (100.0 + column + offset * 0.1)
            if offset == 1 and symbol == "PAAAM00":
                value = None
            worksheet.cell(row, column, value)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_excel_parser_validates_dates_units_and_preserves_zero():
    observations, summary = parse_market_workbook(workbook_bytes())

    forcados = [item for item in observations if item.instrument_key == "forcados"]
    assert forcados[-1].value == 0
    assert forcados[-1].assessment_date == date(2026, 9, 10)
    assert summary["rows_valid"] == 9
    assert set(summary["instruments_affected"]) == {"forcados", "naphtha", "gasoil", "gasoline", "jet"}


def test_excel_parser_rejects_missing_sheet_and_invalid_date():
    with pytest.raises(ExcelValidationError, match="Required sheet"):
        parse_market_workbook(workbook_bytes(include_sheet=False))

    observations, summary = parse_market_workbook(workbook_bytes(invalid_date=True))
    assert len(observations) == 4
    assert summary["rows_rejected"] > 0

    workbook = Workbook()
    workbook.active.title = "Market Risk"
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(ExcelValidationError, match="Required Platts columns"):
        parse_market_workbook(output.getvalue())


def test_excel_history_can_supply_a_full_90_day_calculation_window():
    observations, _ = parse_market_workbook(workbook_bytes(days=90))
    forcados = [
        {"assessment_date": item.assessment_date, "value": item.value}
        for item in observations
        if item.instrument_key == "forcados"
    ]
    stats = calculate_instrument_stats(forcados, window_days=90, min_history_points=60)
    assert stats["window_count"] == 90
    assert stats["history_status"] == "valid"
    assert stats["z_score"] is not None


def test_source_resolution_never_falls_back_across_benchmarks(db_session):
    instrument = MarketInstrument(
        instrument_key="brent",
        display_name="Dated Brent",
        unit="USD/bbl",
        category="crude",
        provider="oilpriceapi",
        provider_symbol="BRENT_CRUDE_USD",
        enabled=True,
    )
    internal = MarketObservation(
        instrument=instrument,
        assessment_date=date(2026, 9, 10),
        value=100,
        unit="USD/BBL",
        provider="internal_excel",
        provider_symbol="PCAAS00",
        retrieved_at=datetime.now(timezone.utc),
    )
    selected, policy, status = resolve_observations(instrument, [internal])
    assert selected == []
    assert policy.preferred_provider == "oilpriceapi"
    assert status == "preferred_unavailable"


def test_source_resolution_prefers_configured_api_source(db_session):
    instrument = MarketInstrument(
        instrument_key="brent",
        display_name="Dated Brent",
        unit="USD/bbl",
        category="crude",
        provider="oilpriceapi",
        provider_symbol="BRENT_CRUDE_USD",
        enabled=True,
    )
    excel_proxy = MarketObservation(
        instrument=instrument,
        assessment_date=date(2026, 9, 9),
        value=99,
        unit="USD/BBL",
        provider="internal_excel",
        provider_symbol="PCAAS00",
        retrieved_at=datetime.now(timezone.utc),
    )
    api_value = MarketObservation(
        instrument=instrument,
        assessment_date=date(2026, 9, 10),
        value=100,
        unit="USD/BBL",
        provider="oilpriceapi",
        provider_symbol="BRENT_CRUDE_USD",
        retrieved_at=datetime.now(timezone.utc),
    )
    selected, policy, status = resolve_observations(instrument, [excel_proxy, api_value])
    assert [item.value for item in selected] == [100]
    assert policy.preferred_provider == "oilpriceapi"
    assert status == "preferred_available"


def test_excel_import_resolves_forcados_as_excel_preferred(client, seed_test_data, db_session):
    files = {"workbook": ("hybrid.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    response = client.post("/api/market/import-excel", files=files)

    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["rows_stored"] == 9
    assert payload["rows_rejected"] == 1
    assert set(payload["instruments_affected"]) == {"forcados", "naphtha", "gasoil", "gasoline", "jet"}

    # No second workbook-current copy is written to the API observation table.
    assert db_session.scalars(select(MarketObservation).where(MarketObservation.provider == "internal_excel")).all() == []

    snapshot = client.get("/api/dashboard/snapshot")
    assert snapshot.status_code == 200
    stats = {item["instrument_id"]: item for item in snapshot.json()["market_stats"]}
    assert stats["forcados"]["provider"] == "platts_excel"
    assert stats["forcados"]["provider_symbol"] == "PCABC00"
    assert stats["forcados"]["latest_date"] == "2026-09-10"
    assert stats["forcados"]["current_value"] is not None

    assert {item["instrument_id"] for item in snapshot.json()["product_spreads"]} == {
        "naphtha", "gasoil", "gasoline", "jet"
    }


def test_excel_import_is_idempotent_and_selected_in_history(client, seed_test_data):
    files = {"workbook": ("hybrid.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    first = client.post("/api/market/import-excel", files=files)
    assert first.status_code == 200, first.text
    first_data = first.json()
    assert first_data["provider"] == "platts_excel"
    assert first_data["rows_stored"] == 9

    second = client.post("/api/market/import-excel", files=files)
    assert second.status_code == 200, second.text
    assert second.json()["rows_updated"] == 9
    assert second.json()["rows_stored"] == 0

    history = client.get("/api/market/history?instrument=forcados&days=90")
    assert history.status_code == 200
    assert history.json()["observations"][0]["provider"] == "platts_excel"
    assert history.json()["observations"][0]["sourceType"] == "excel"


def test_gemini_context_uses_resolved_excel_source(client, seed_test_data, db_session):
    response = client.post(
        "/api/market/import-excel",
        files={"workbook": ("hybrid.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200

    context = build_dashboard_ai_context(db_session)
    forcados = next(item for item in context["market_stats"] if item["instrument_id"] == "forcados")
    assert forcados["provider"] == "platts_excel"
    assert forcados["provider_symbol"] == "PCABC00"
    assert forcados["benchmark_definition"] == "Platts Forcados FOB Nigeria"
    assert forcados["benchmark_status"] == "confirmed"

    # Refined products share the exact workbook lineage.
    naphtha = next(item for item in context["market_stats"] if item["instrument_id"] == "naphtha")
    assert naphtha["benchmark_status"] == "confirmed"
    assert naphtha["source_provider"] == "platts_excel"
    assert naphtha["source_symbol"] == "PAAAM00"


def test_official_source_policy_is_platts_confirmed_for_all_seven():
    from app.services.source_resolution import SOURCE_POLICIES
    assert SOURCE_POLICIES["brent"].benchmark_status == "confirmed"
    assert SOURCE_POLICIES["wti"].benchmark_status == "confirmed"
    assert SOURCE_POLICIES["forcados"].benchmark_status == "confirmed"
    assert SOURCE_POLICIES["naphtha"].benchmark_status == "confirmed"
    assert SOURCE_POLICIES["gasoil"].benchmark_status == "confirmed"
    assert SOURCE_POLICIES["gasoline"].benchmark_status == "confirmed"
    assert SOURCE_POLICIES["jet"].benchmark_status == "confirmed"
