from datetime import date, datetime, timezone
from pathlib import Path
from io import BytesIO
import pytest
from openpyxl import Workbook
from sqlalchemy import select, func
from app.api.market import observation_to_normalized
from app.models.market import MarketInstrument, MarketObservation
from app.models.historical import HistoricalMarketObservation
from app.services.excel_ingestion import parse_market_workbook
from app.services.historical_market import HISTORICAL_MAPPING, historical_analytics


@pytest.mark.parametrize("unit,value,key,expected,factor", [
    ("USD/mt", 731.37, "gasoil", 731.37/7.44, 7.44),
    ("tonne", 731.37, "gasoil", 731.37/7.44, 7.44),
    ("metric_ton", 890, "naphtha", 100, 8.9),
    ("USD/gallon", 2, "gasoline", 84, None),
    ("gallon", 2, "jet", 84, None),
    ("USD/bbl", 84, "jet", 84, None),
    ("barrel", 84, "brent", 84, None),
    ("USD/gallon", 0, "jet", 0, None),
    ("unknown", 84, "jet", None, None),
])
def test_api_normalization_routes_native_units(unit,value,key,expected,factor):
    obs=MarketObservation(value=value,unit=unit,provider="oilpriceapi",provider_symbol="TEST",
        assessment_date=date(2026,9,10),retrieved_at=datetime.now(timezone.utc))
    normalized=observation_to_normalized(obs,key)
    assert normalized.value==value
    assert normalized.converted_value==pytest.approx(expected) if expected is not None else normalized.converted_value is None
    assert normalized.barrels_per_mt==factor


def sample_workbook():
    workbook=Workbook(); ws=workbook.active; ws.title="Market Risk"; ws.cell(21,6,"AssessDate")
    for col,(key,(symbol,label,unit,*_)) in enumerate(HISTORICAL_MAPPING.items(),7):
        ws.cell(17,col,symbol);ws.cell(18,col,label);ws.cell(19,col,unit.split("/")[1]);ws.cell(20,col,"USD")
        for row,dt in enumerate([date(2026,9,10),date(2026,8,12),date(2026,8,11),date(2026,6,13),date(2026,6,12)],22):
            ws.cell(row,6,datetime.combine(dt,datetime.min.time()))
            ws.cell(row,col,100 if key=="brent" else 779 if key=="jet" else 200)
    # An explicitly missing value remains missing; valid zero remains present.
    ws.cell(22,8).value=None
    ws.cell(23,8).value=0
    sheet=workbook.create_sheet("PCAAS00_minus_PJAAV00_jet_crack")
    sheet.append(["Date","Dated Brent PCAAS00_USD_per_BBL","Jet FOB NWE Cargo PJAAV00_USD_per_MT"])
    for offset in range(6):
        sheet.append([datetime(2026,9,10-offset),100,779])
    # Duplicate raw rows must not increase stored observation counts.
    sheet.append([datetime(2026,9,10),100,999])
    output=BytesIO();workbook.save(output);return output.getvalue()


def test_historical_import_is_separate_idempotent_and_preserves_provenance(client,seed_test_data,db_session):
    before=client.get("/api/market/latest").json()
    payload=sample_workbook()
    first=client.post("/api/market/import-excel",files={"workbook":("history.xlsx",payload)})
    assert first.status_code==200, first.text
    assert set(first.json()["instruments_affected"])==set(HISTORICAL_MAPPING)
    assert first.json()["rows_rejected"]>=2
    result=client.get("/api/market/historical").json()
    series={s["instrument"]:s for s in result["instruments"]}
    assert series["jet"]["observation_count"]==10
    assert series["jet"]["spread"]==pytest.approx(779/7.7892-100)
    assert series["jet"]["spread"]>0
    assert series["naphtha"]["latest_date"]=="2026-08-12"
    assert series["naphtha"]["latest"]==0
    assert series["brent"]["count_30d"]==2
    assert series["brent"]["count_90d"]==4
    assert series["brent"]["freshness"]=="stale"
    assert series["brent"]["change_30d"]==0
    rows=db_session.scalars(select(HistoricalMarketObservation)).all()
    assert all(r.provider=="platts_excel" and r.workbook_source=="history.xlsx" and len(r.workbook_sha256)==64 and r.imported_at for r in rows)
    after=client.get("/api/market/latest").json()
    exclude_forcados=lambda v:[o for o in v["observations"] if o["instrumentId"]!="forcados"]
    # Freshness ages naturally advance, so compare identity/value rather than wall-clock age.
    identity=lambda v:[(o["instrumentId"],o["providerSymbol"],o["value"],o["assessmentDate"]) for o in exclude_forcados(v)]
    assert identity(before)==identity(after)
    count=len(rows)
    second=client.post("/api/market/import-excel",files={"workbook":("history.xlsx",payload)})
    assert second.status_code==200
    assert second.json()["rows_stored"]==0
    assert second.json()["rows_updated"]==count
    assert db_session.scalar(select(func.count()).select_from(HistoricalMarketObservation))==count


def test_empty_history_is_unavailable_not_zero(client):
    response=client.get("/api/market/historical")
    assert response.status_code==200
    assert len(response.json()["instruments"])==7
    assert all(s["latest"] is None and s["freshness"]=="unavailable" for s in response.json()["instruments"])


@pytest.mark.skipif(not Path("uploads/active-market-data.xlsx").exists(),reason="Controlled local audit workbook not bundled")
def test_audited_workbook_longer_copies_selected_once():
    observations,summary=parse_market_workbook(Path("uploads/active-market-data.xlsx").read_bytes())
    assert summary["rows_valid"]==2488
    for key in HISTORICAL_MAPPING:
        rows=[o for o in observations if o.instrument_key==key]
        assert len(rows)==(175 if key=="forcados" else 173 if key=="wti" else 428)
        assert len({o.assessment_date for o in rows})==len(rows)
        assert rows[-1].assessment_date==date(2026,9,10)
        assert rows[-1].conversion_factor==HISTORICAL_MAPPING[key][3]
        if rows[-1].conversion_factor:
            assert rows[-1].converted_value==pytest.approx(rows[-1].value/rows[-1].conversion_factor)


def test_refresh_materialized_results_retains_history_and_api_series(client, seed_test_data, db_session):
    from app.services.historical_market import persist_history
    from hashlib import sha256
    payload = sample_workbook()
    original, _ = parse_market_workbook(payload)
    from dataclasses import replace
    original.append(replace(next(o for o in original if o.instrument_key == "brent"), assessment_date=date(2025, 1, 2)))
    # Seed older observations, including history outside the materialized table.
    persist_history(db_session, original, payload, "baseline.xlsx", datetime.now(timezone.utc))
    db_session.commit()
    before_count = db_session.scalar(select(func.count()).select_from(HistoricalMarketObservation))
    baseline = {s["instrument"]: s for s in client.get("/api/market/historical").json()["instruments"]}
    assert baseline["brent"]["latest_date"] == "2026-09-10"
    api_before = [(r.id, r.provider, r.provider_symbol, r.value, r.assessment_date)
                  for r in db_session.scalars(select(MarketObservation)).all() if r.provider != "internal_excel"]
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(payload)); market = workbook["Market Risk"]
    for col, (key, mapping) in enumerate(HISTORICAL_MAPPING.items(), 7):
        market.cell(22, 6, datetime(2026, 10, 14))
        market.cell(22, col, 100 if key == "brent" else 900)
    # A null on a valid assessment must not fabricate an observation or zero.
    market.cell(27, 6, datetime(2026, 10, 13))
    market.cell(27, 8).value = None
    # Invalid date with numeric values must not produce observations.
    market.cell(28, 6, "not a date"); market.cell(28, 8, 123)
    output = BytesIO(); workbook.save(output); refreshed = output.getvalue()
    response = client.post("/api/market/import-excel", files={"workbook":("refreshed.xlsx", refreshed)})
    assert response.status_code == 200, response.text
    assert response.json()["historical_rows_inserted"] == 7
    assert response.json()["latest_historical_assessment_date"] == "2026-10-14"
    assert set(response.json()["mapped_historical_instruments"]) == set(HISTORICAL_MAPPING)
    series = {s["instrument"]: s for s in client.get("/api/market/historical").json()["instruments"]}
    assert all(s["latest_date"] == "2026-10-14" for s in series.values())
    for key, (symbol, _, unit, factor, *_) in HISTORICAL_MAPPING.items():
        latest = db_session.scalar(select(HistoricalMarketObservation).where(
            HistoricalMarketObservation.instrument_key == key,
            HistoricalMarketObservation.assessment_date == date(2026, 10, 14)))
        assert latest.provider_symbol == symbol and latest.native_unit == unit
        assert latest.worksheet == "Market Risk"
        assert latest.workbook_sha256 == sha256(refreshed).hexdigest()
        assert float(latest.conversion_factor) == pytest.approx(factor) if factor else latest.conversion_factor is None
        if factor:
            assert float(latest.converted_value) == pytest.approx(900 / factor)
            assert series[key]["spread"] == pytest.approx(900 / factor - 100)
    after_count = db_session.scalar(select(func.count()).select_from(HistoricalMarketObservation))
    assert after_count == before_count + 7
    assert db_session.scalar(select(HistoricalMarketObservation).where(
        HistoricalMarketObservation.instrument_key == "brent",
        HistoricalMarketObservation.assessment_date == date(2025, 1, 2))) is not None
    assert db_session.scalar(select(HistoricalMarketObservation).where(
        HistoricalMarketObservation.instrument_key == "jet",
        HistoricalMarketObservation.assessment_date == date(2026, 9, 10))) is not None
    assert db_session.scalar(select(HistoricalMarketObservation).where(
        HistoricalMarketObservation.instrument_key == "naphtha",
        HistoricalMarketObservation.assessment_date == date(2026, 10, 13))) is None
    repeat = client.post("/api/market/import-excel", files={"workbook":("refreshed.xlsx", refreshed)}).json()
    assert repeat["historical_rows_inserted"] == 0
    assert repeat["historical_rows_updated_or_unchanged"] == response.json()["rows_valid"]
    assert db_session.scalar(select(func.count()).select_from(HistoricalMarketObservation)) == after_count
    assert api_before == [(r.id, r.provider, r.provider_symbol, r.value, r.assessment_date)
        for r in db_session.scalars(select(MarketObservation)).all() if r.provider != "internal_excel"]


def test_unmapped_daily_template_cannot_populate_history(client):
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Market Risk"
    sheet.append(["AssessDate", "UNRELATED_SYMBOL"]); sheet.append([datetime(2026, 10, 14), 12])
    output = BytesIO(); workbook.save(output)
    response = client.post("/api/market/import-excel", files={"workbook":("Daily_Historical_Template1.xlsx", output.getvalue())})
    assert response.status_code == 400
    assert all(s["observation_count"] == 0 for s in client.get("/api/market/historical").json()["instruments"])
