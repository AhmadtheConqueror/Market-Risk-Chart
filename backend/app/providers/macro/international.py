"""Official USA/global publications. Source structures documented in MACRO_SOURCE_DISCOVERY.md."""
from __future__ import annotations

import csv
import io
import re
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from urllib.parse import urljoin, urlparse

import httpx
from lxml import html
from openpyxl import load_workbook
from pypdf import PdfReader

from app.providers.macro.base import MacroDataProvider, MacroObservation

BLS = "U.S. Bureau of Labor Statistics"
FED = "Federal Reserve / FRED"
BEA = "U.S. Bureau of Economic Analysis"
EIA = "U.S. Energy Information Administration"
IMF = "IMF World Economic Outlook"
WB = "World Bank Commodity Price Data / Pink Sheet"
CHINA = "National Bureau of Statistics of China"

# One provider per indicator isolates failures even within a shared publisher.
SPECS = {
    "usa_headline_inflation": ("Headline Inflation", "%", BLS, "monthly", "CUUR0000SA0"),
    "usa_core_inflation": ("Core Inflation", "%", BLS, "monthly", "CUUR0000SA0L1E"),
    "usa_real_gdp_growth": ("Real GDP Growth", "%", BEA, "quarterly", "GDPhistQ"),
    "usa_policy_rate": ("Effective Federal Funds Rate", "%", FED, "daily", "EFFR"),
    "usa_unemployment_rate": ("Unemployment Rate", "%", BLS, "monthly", "LNS14000000"),
    "usa_crude_inventories": ("Commercial Crude Inventories", "million barrels", EIA, "weekly", "WCESTUS1"),
    "global_real_gdp_growth": ("World Real GDP Growth", "%", IMF, "weo", "NGDP_RPCH"),
    "global_inflation": ("World Inflation", "%", IMF, "weo", "PCPIPCH"),
    "china_manufacturing_pmi": ("China Manufacturing PMI", "index", CHINA, "monthly", "manufacturing PMI"),
    "global_oil_demand_growth": ("Global Oil Demand Growth", "mbpd", EIA, "outlook", "STEO Table 3e"),
    "global_energy_price_index": ("Energy Price Index", "index", WB, "monthly", "Monthly Indices / Energy"),
    "broad_usd_index": ("Broad U.S. Dollar Index", "index", FED, "daily", "DTWEXBGS"),
}


def number(value) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        result = Decimal(str(value).strip().replace(",", ""))
        return result if result.is_finite() else None
    except InvalidOperation:
        return None


def utc_date(text, pattern="%Y-%m-%d"):
    return datetime.strptime(text.strip(), pattern).replace(tzinfo=timezone.utc)


def observation(key, value, period, url, published=None, **metadata):
    title, unit, source, cycle, series = SPECS[key]
    return MacroObservation(key, title, value, unit, period, source, url, published,
        metadata_json={"geography": "usa" if key.startswith("usa_") else "global",
                       "expected_publication_cycle": cycle, "series_id": series, **metadata})


def parse_bls(payload, key, url):
    series_id = SPECS[key][4]
    if payload.get("status") != "REQUEST_SUCCEEDED":
        raise ValueError("BLS request did not succeed")
    series = next((s for s in payload.get("Results", {}).get("series", []) if s.get("seriesID") == series_id), None)
    if not series:
        raise ValueError("BLS canonical series missing")
    values = {}
    for row in series.get("data", []):
        if re.fullmatch(r"M(?:0[1-9]|1[0-2])", row.get("period", "")):
            value = number(row.get("value"))
            if value is not None:
                values[(int(row["year"]), int(row["period"][1:]))] = value
    result = []
    for (year, month), value in sorted(values.items()):
        meta = {"definition": "Seasonally adjusted civilian unemployment rate"}
        if key != "usa_unemployment_rate":
            prior = values.get((year - 1, month))
            if prior is None or prior <= 0:
                continue
            meta = {"definition": "CPI year-on-year percent change, not seasonally adjusted" if key == "usa_headline_inflation" else "CPI excluding food and energy, year-on-year percent change, not seasonally adjusted",
                    "source_unit": "CPI index", "current_index": str(value), "year_earlier_index": str(prior),
                    "transformation": "(current index / matched year-earlier index - 1) * 100"}
            value = (value / prior - 1) * 100
        period = datetime(year, month, 1).strftime("%b %Y")
        result.append(observation(key, value, period, url, **meta))
    if not result:
        raise ValueError("No complete BLS observations; no missing release is interpolated")
    return result


def parse_fred(text, key, url):
    series = SPECS[key][4]
    rows = csv.DictReader(io.StringIO(text))
    if not rows.fieldnames or "observation_date" not in rows.fieldnames or series not in rows.fieldnames:
        raise ValueError("FRED CSV columns changed")
    result = []
    for row in rows:
        value = number(row.get(series))
        if value is None:
            continue
        date = utc_date(row["observation_date"])
        result.append(observation(key, value, date.strftime("%Y-%m-%d"), url,
            definition="Effective Federal Funds Rate, percent, not seasonally adjusted" if series == "EFFR" else "Nominal broad trade-weighted US dollar index, January 2006=100, not seasonally adjusted",
            index_base="January 2006=100" if series == "DTWEXBGS" else None))
    if not result:
        raise ValueError("FRED publication has no numeric observations")
    return result


def parse_bea(content, url):
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        rows = list(book["GDPhistQ"].values)
        if "Seasonally adjusted at annual rates" not in str(rows[2][0]) or "Real Gross Domestic Product" not in str(rows[1][0]):
            raise ValueError("BEA GDP annual-rate definition changed")
        pub = utc_date(str(rows[0][0]), "%B %d, %Y")
        latest = re.search(r"(\d{4})Q([1-4])", str(rows[1][0]))
        gdp = next(r for r in rows if r[0] == "Gross domestic product (GDP)")
        if not latest:
            raise ValueError("BEA quarter absent")
        periods = {f"{latest[1]}Q{latest[2]}": gdp[1]}
        # These are exact historic GDP comparisons, not adjacent unrelated rows.
        for position in (2, 4, 6):
            if re.fullmatch(r"\d{4}Q[1-4]", str(gdp[position])):
                periods[str(gdp[position])] = gdp[position + 1]
        return [observation("usa_real_gdp_growth", number(value), f"Q{period[-1]} {period[:4]}", url, pub,
            definition="Quarterly real GDP percent change from preceding quarter, seasonally adjusted annual rate",
            source_table="GDPhistQ; Gross domestic product (GDP)") for period, value in sorted(periods.items()) if number(value) is not None]
    finally:
        book.close()


def parse_eia_stocks(content, url):
    tree = html.fromstring(content)
    title = " ".join(tree.xpath("//title/text()"))
    if "Ending Stocks excluding SPR of Crude Oil" not in title or "Thousand Barrels" not in title:
        raise ValueError("EIA commercial crude excluding SPR definition changed")
    release = re.search(r"Release Date:\s*(\d{1,2}/\d{1,2}/\d{4})", tree.text_content())
    pub = utc_date(release[1], "%m/%d/%Y") if release else None
    result = []
    for row in tree.xpath("//tr[td[@class='B6']]"):
        cells = row.xpath("./td")
        year_match = re.search(r"(\d{4})-[A-Za-z]{3}", cells[0].text_content())
        if not year_match:
            continue
        for i in range(1, len(cells) - 1, 2):
            period = cells[i].text_content().strip()
            value = number(cells[i + 1].text_content())
            if not re.fullmatch(r"\d{2}/\d{2}", period) or value is None:
                continue
            date = utc_date(f"{year_match[1]}-{period.replace('/', '-')}")
            result.append(observation("usa_crude_inventories", value / 1000, date.strftime("%Y-%m-%d"), url, pub,
                definition="US commercial ending stocks of crude oil excluding Strategic Petroleum Reserve, week ending",
                source_unit="thousand barrels", conversion="source / 1000"))
    if not result:
        raise ValueError("EIA stock publication has no verified weeks")
    return result


def parse_imf(payload, metadata, key, url):
    series = SPECS[key][4]
    info = metadata.get("indicators", {}).get(series, {})
    if info.get("dataset") != "WEO" or info.get("unit") != "Annual percent change":
        raise ValueError("IMF WEO series definition changed")
    pub = utc_date(info["last-modified"], "%Y-%m-%d %H:%M:%S")
    world = payload.get("values", {}).get(series, {}).get("WEOWORLD")
    if not world:
        raise ValueError("IMF WEOWORLD aggregate missing; country values are not substitutes")
    result = []
    for year, raw in sorted(world.items()):
        value = number(raw)
        if not re.fullmatch(r"\d{4}", year) or value is None:
            continue
        result.append(observation(key, value, year, url, pub,
            definition=info["label"] + ", World aggregate, annual percent change",
            reference_year=int(year), forecast=int(year) >= pub.year,
            publication_edition=info["source"], publication_date_basis="IMF official series last-modified",
            source_unit=info["unit"], aggregate="WEOWORLD"))
    if not result:
        raise ValueError("No IMF world observations")
    return result


def parse_china(content, url):
    tree = html.fromstring(content)
    text = re.sub(r"\s+", "", tree.text_content())
    # Exact manufacturing paragraph avoids non-manufacturing/business activity PMI.
    match = re.search(r"(\d{4})年(\d{1,2})月中国采购经理指数运行情况", text)
    value = re.search(r"(?<!非)制造业采购经理指数(?:（PMI）|\(PMI\))?为([0-9]+(?:\.[0-9]+)?)%", text)
    if not match or not value:
        raise ValueError("China NBS official manufacturing PMI paragraph missing")
    date_match = re.search(r"/t(\d{8})_", url)
    pub = utc_date(date_match[1], "%Y%m%d") if date_match else None
    values = {(int(match[1]), int(match[2])): Decimal(value[1])}
    tables = [table for table in tree.xpath("//table")
              if any(cell.text_content().strip() == "PMI" for cell in table.xpath(".//td|.//th"))]
    for table in tables:
        for row in table.xpath(".//tr"):
            cells = row.xpath("./td|./th")
            period = re.fullmatch(r"(\d{4})年(\d{1,2})月", re.sub(r"\s+", "", cells[0].text_content())) if cells else None
            numeric = number(cells[1].text_content()) if len(cells) > 1 else None
            if period and numeric is not None:
                pair = (int(period[1]), int(period[2]))
                if pair in values and values[pair] != numeric:
                    raise ValueError("China manufacturing paragraph and table disagree")
                values[pair] = numeric
    result = []
    for (year, month), numeric in sorted(values.items()):
        previous = values.get((year-1, 12) if month == 1 else (year, month-1))
        metadata = {"definition": "Official China NBS manufacturing purchasing managers index; index scale, not percent change"}
        if previous is not None:
            change = numeric - previous
            metadata.update(change=float(change), change_unit="index points",
                direction="up" if change > 0 else "down" if change < 0 else "unchanged")
        result.append(observation("china_manufacturing_pmi", numeric, datetime(year, month, 1).strftime("%b %Y"), url, pub, **metadata))
    return result



def parse_world_bank(content, url):
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        rows = list(book["Monthly Indices"].values)
        if not any("2010=100" in str(r[0]) for r in rows[:10]):
            raise ValueError("World Bank energy index base changed")
        headers = next(r for r in rows[:12] if "Energy" in r)
        column = headers.index("Energy")
        updated = next(re.search(r"Updated on (.+)", str(r[0])) for r in rows[:10] if "Updated on " in str(r[0]))
        pub = utc_date(updated[1], "%B %d, %Y")
        result = []
        for row in rows:
            match = re.fullmatch(r"(\d{4})M(\d{2})", str(row[0]))
            value = number(row[column])
            if match and value is not None:
                period = datetime(int(match[1]), int(match[2]), 1).strftime("%b %Y")
                result.append(observation("global_energy_price_index", value, period, url, pub,
                    definition="World Bank Pink Sheet Energy price index, nominal US dollars, 2010=100", index_base="2010=100"))
        if not result:
            raise ValueError("World Bank energy index observations missing")
        return result
    finally:
        book.close()


def parse_steo(text, release_text, url):
    if "World Petroleum and Other Liquid Fuels Consumption" not in text or "million barrels per day" not in text.lower():
        raise ValueError("EIA STEO Table 3e identity/unit changed")
    # Annual header has twelve quarterly labels followed by three explicit annual years.
    header = re.search(r"Q1\s+Q2\s+Q3\s+Q4\s+Q1\s+Q2\s+Q3\s+Q4\s+Q1\s+Q2\s+Q3\s+Q4\s+(20\d{2})\s+(20\d{2})\s+(20\d{2})", text)
    world = re.search(r"World total[^\n]*", text)
    release = re.search(r"Release Date:\s*([A-Za-z]+ \d{1,2}, \d{4})", release_text)
    if not header or not world or not release:
        raise ValueError("EIA STEO annual columns/publication date missing")
    values = re.findall(r"[+-]?\d+\.\d+", world[0])
    if len(values) != 15:
        raise ValueError("EIA STEO world consumption column count changed")
    pub = utc_date(release[1], "%B %d, %Y")
    annual = dict(zip(map(int, header.groups()), map(Decimal, values[-3:])))
    return [observation("global_oil_demand_growth", annual[year] - annual[year-1], str(year), url, pub,
        definition="Year-on-year change in annual world petroleum and other liquid fuels consumption, million barrels per day",
        transformation="current reference-year annual world consumption minus preceding-year consumption",
        current_consumption_mbpd=str(annual[year]), previous_consumption_mbpd=str(annual[year-1]),
        reference_year=year, forecast=year >= pub.year, publication_edition="STEO " + pub.strftime("%B %Y"))
        for year in sorted(annual) if year-1 in annual]


class InternationalMacroProvider(MacroDataProvider):
    def __init__(self, key: str):
        self.key = key

    @property
    def provider_name(self):
        return SPECS[self.key][2] + " / " + self.key

    async def fetch_latest(self):
        key = self.key
        async with httpx.AsyncClient(timeout=25, follow_redirects=True) as client:
            async def get(url):
                response = await client.get(url)
                response.raise_for_status()
                return response

            async def discovered(page, label, host, suffix):
                tree = html.fromstring((await get(page)).content)
                links = [urljoin(page, a.get("href")) for a in tree.xpath("//a[@href]")
                         if label in a.text_content() or label in a.get("href")]
                links = [u for u in links if urlparse(u).hostname == host and urlparse(u).path.endswith(suffix)]
                if not links:
                    raise ValueError("Official publication link not found: " + label)
                return links[0]

            if SPECS[key][2] == BLS:
                url = "https://api.bls.gov/publicAPI/v2/timeseries/data/" + SPECS[key][4]
                return parse_bls((await get(url)).json(), key, url)
            if SPECS[key][2] == FED:
                url = "https://fred.stlouisfed.org/graph/fredgraph.csv?id=" + SPECS[key][4]
                return parse_fred((await get(url)).text, key, url)
            if key == "usa_real_gdp_growth":
                url = await discovered("https://www.bea.gov/data/gdp/gross-domestic-product", "Historical Comparisons", "www.bea.gov", ".xlsx")
                return parse_bea((await get(url)).content, url)
            if key == "usa_crude_inventories":
                url = "https://www.eia.gov/dnav/pet/hist/LeafHandler.ashx?n=PET&s=WCESTUS1&f=W"
                return parse_eia_stocks((await get(url)).content, url)
            if SPECS[key][2] == IMF:
                url = "https://www.imf.org/external/datamapper/api/v1/" + SPECS[key][4] + "/WEOWORLD"
                metadata = (await get("https://www.imf.org/external/datamapper/api/v1/indicators")).json()
                return parse_imf((await get(url)).json(), metadata, key, url)
            if key == "china_manufacturing_pmi":
                # The release list uses zxfb aliases; official canonical pages use zxfbhjd.
                page = "https://www.stats.gov.cn/sj/zxfb/"
                tree = html.fromstring((await get(page)).content.decode("utf-8"))
                candidates = [urljoin(page, a.get("href")) for a in tree.xpath("//a[@href]") if "中国采购经理指数运行情况" in a.text_content()]
                candidates = [u for u in candidates if urlparse(u).hostname == "www.stats.gov.cn" and re.search(r"/t\d{8}_", u)]
                if not candidates:
                    raise ValueError("China NBS purchasing managers release missing")
                url = max(candidates, key=lambda u: re.search(r"/t(\d{8})_", u)[1]).replace("/sj/zxfb/", "/sj/zxfbhjd/")
                return parse_china((await get(url)).content.decode("utf-8"), url)
            if key == "global_energy_price_index":
                url = await discovered("https://www.worldbank.org/en/research/commodity-markets", "CMO-Historical-Data-Monthly.xlsx", "thedocs.worldbank.org", ".xlsx")
                return parse_world_bank((await get(url)).content, url)
            if key == "global_oil_demand_growth":
                url = "https://www.eia.gov/outlooks/steo/tables/pdf/3etab.pdf"
                release = html.fromstring((await get("https://www.eia.gov/outlooks/steo/")).content).text_content()
                text = "\n".join(p.extract_text(extraction_mode="layout") for p in PdfReader(io.BytesIO((await get(url)).content)).pages)
                return parse_steo(text, release, url)
        raise ValueError("Unsupported international indicator")


def get_international_macro_providers():
    return [InternationalMacroProvider(key) for key in SPECS]
