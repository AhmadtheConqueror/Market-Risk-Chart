from __future__ import annotations

import math
from dataclasses import dataclass
from datetime import date, datetime, time, timezone
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from app.services.source_resolution import SOURCE_POLICIES


EXCEL_PROVIDER = "internal_excel"
EXCEL_SHEET = "Market Risk"
SYMBOL_TO_INSTRUMENT = {policy.preferred_symbol: key for key, policy in SOURCE_POLICIES.items() if policy.preferred_provider == EXCEL_PROVIDER}
EXPECTED_UNITS = {
    "PCABC00": "USD/BBL",
    "PAAAM00": "USD/MT",
    "AAVJI00": "USD/MT",
    "PGABM00": "USD/MT",
    "PJAAV00": "USD/MT",
}


class ExcelValidationError(ValueError):
    pass


@dataclass(frozen=True)
class ExcelObservation:
    instrument_key: str
    provider_symbol: str
    value: float
    unit: str
    assessment_date: date
    benchmark_definition: str


def parse_excel_date(value: Any, epoch: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        try:
            converted = from_excel(value, epoch)
            return converted.date() if isinstance(converted, datetime) else converted
        except (TypeError, ValueError, OverflowError):
            return None
    if isinstance(value, str) and value.strip():
        raw = value.strip()
        for parser in (datetime.fromisoformat,):
            try:
                return parser(raw.replace("Z", "+00:00")).date()
            except ValueError:
                pass
        for fmt in ("%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(raw, fmt).date()
            except ValueError:
                pass
    return None


def parse_nullable_number(value: Any) -> float | None:
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def parse_market_workbook(file_bytes: bytes) -> tuple[list[ExcelObservation], dict[str, Any]]:
    try:
        # The source workbook is modest in size. Normal mode keeps random cell
        # access efficient while preserving Excel date values via openpyxl.
        workbook = load_workbook(BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        raise ExcelValidationError("Workbook could not be opened as a valid .xlsx file.") from exc

    if EXCEL_SHEET not in workbook.sheetnames:
        raise ExcelValidationError(f"Required sheet '{EXCEL_SHEET}' is missing.")

    worksheet = workbook[EXCEL_SHEET]
    symbol_row, symbol_columns = _find_symbol_columns(worksheet)
    if not symbol_columns:
        raise ExcelValidationError("Required Platts columns were not found in the Market Risk sheet.")

    date_column = _find_date_column(worksheet, symbol_row)
    if date_column is None:
        raise ExcelValidationError("Required AssessDate column was not found in the Market Risk sheet.")

    unit_row = symbol_row + 2
    currency_row = symbol_row + 3
    required_units: dict[str, str] = {}
    for symbol, column in symbol_columns.items():
        commodity_unit = _clean_text(worksheet.cell(unit_row, column).value).upper()
        currency = _clean_text(worksheet.cell(currency_row, column).value).upper()
        unit = f"{currency}/{commodity_unit}" if currency and commodity_unit else ""
        expected = EXPECTED_UNITS[symbol]
        if unit != expected:
            raise ExcelValidationError(
                f"Unexpected unit for {symbol}: found '{unit or 'blank'}', expected '{expected}'."
            )
        required_units[symbol] = unit

    observations: list[ExcelObservation] = []
    seen: set[tuple[str, date]] = set()
    errors: list[str] = []
    rows_read = 0
    rows_rejected = 0
    upload_date = datetime.now(timezone.utc)

    for row_number in range(symbol_row + 5, (worksheet.max_row or symbol_row) + 1):
        date_value = worksheet.cell(row_number, date_column).value
        prices = {symbol: worksheet.cell(row_number, column).value for symbol, column in symbol_columns.items()}
        if date_value is None and not any(value is not None for value in prices.values()):
            continue
        rows_read += 1
        assessment_date = parse_excel_date(date_value, workbook.epoch)
        if assessment_date is None:
            rows_rejected += 1
            _record_error(errors, f"Row {row_number}: invalid assessment date.")
            continue

        for symbol, raw_price in prices.items():
            value = parse_nullable_number(raw_price)
            if value is None:
                rows_rejected += 1
                if raw_price not in (None, ""):
                    _record_error(errors, f"Row {row_number}, {symbol}: price is not numeric.")
                continue
            key = (symbol, assessment_date)
            if key in seen:
                rows_rejected += 1
                _record_error(errors, f"Row {row_number}, {symbol}: duplicate assessment date.")
                continue
            seen.add(key)
            instrument_key = SYMBOL_TO_INSTRUMENT[symbol]
            observations.append(
                ExcelObservation(
                    instrument_key=instrument_key,
                    provider_symbol=symbol,
                    value=value,
                    unit=required_units[symbol],
                    assessment_date=assessment_date,
                    benchmark_definition=SOURCE_POLICIES[instrument_key].benchmark_definition,
                )
            )

    if not observations:
        raise ExcelValidationError("No valid supported market observations were found in the workbook.")

    dates = [item.assessment_date for item in observations]
    summary = {
        "rows_read": rows_read,
        "rows_valid": len(observations),
        "rows_rejected": rows_rejected,
        "instruments_affected": sorted({item.instrument_key for item in observations}),
        "date_range": {"from": min(dates).isoformat(), "to": max(dates).isoformat()},
        "uploaded_at": upload_date.isoformat(),
        "errors": errors,
    }
    return observations, summary


def _find_symbol_columns(worksheet: Any) -> tuple[int, dict[str, int]]:
    best_row = 0
    best_columns: dict[str, int] = {}
    for row_number in range(1, min(60, worksheet.max_row or 60) + 1):
        columns: dict[str, int] = {}
        for column in range(1, (worksheet.max_column or 1) + 1):
            value = _clean_text(worksheet.cell(row_number, column).value).upper()
            if value in SYMBOL_TO_INSTRUMENT:
                columns[value] = column
        if len(columns) > len(best_columns):
            best_row, best_columns = row_number, columns
    return best_row, best_columns


def _find_date_column(worksheet: Any, symbol_row: int) -> int | None:
    for row_number in range(symbol_row, min(symbol_row + 8, worksheet.max_row or symbol_row) + 1):
        for column in range(1, (worksheet.max_column or 1) + 1):
            value = _clean_text(worksheet.cell(row_number, column).value).lower().replace(" ", "")
            if value in {"assessdate", "assessmentdate", "date", "timestamp"}:
                return column
    return None


def _clean_text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _record_error(errors: list[str], message: str) -> None:
    if len(errors) < 50:
        errors.append(message)
