"""Bounded deterministic extraction with real source locators; never executes uploads."""
import csv
import json
from io import BytesIO, StringIO
from pathlib import Path
import re
from zipfile import ZipFile, BadZipFile
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

EXTENSIONS={".pdf":"PDF",".docx":"DOCX",".xlsx":"XLSX",".csv":"CSV",".txt":"TXT"}
MIME={"PDF":"application/pdf","DOCX":"application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "XLSX":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet","CSV":"text/csv","TXT":"text/plain"}


class ExtractionError(ValueError):
    def __init__(self,message,limitations=None):
        super().__init__(message)
        self.limitations=limitations or []


def safe_filename(filename):
    name=Path(str(filename or "").replace("\\","/")).name
    name=re.sub(r'[\x00-\x1f<>:"/\\|?*]',"_",name).strip(" .")[:255]
    if not name: raise ExtractionError("A valid filename is required.")
    return name


def inspect_zip(content):
    try:
        z=ZipFile(BytesIO(content))
        entries=z.infolist()
        if len(entries)>3000 or sum(i.file_size for i in entries)>60*1024*1024:
            raise ExtractionError("Office archive exceeds the safe extraction limit.")
        if any(i.flag_bits&1 or i.filename.startswith(("/","\\")) or ".." in i.filename.replace("\\","/").split("/") for i in entries):
            raise ExtractionError("Unsafe or encrypted Office archive.")
        if any(i.file_size>25*1024*1024 for i in entries): raise ExtractionError("Office XML part exceeds the safe extraction limit.")
        return z
    except BadZipFile as exc: raise ExtractionError("Invalid Office archive.") from exc


def validate_file(filename,content,mime,max_bytes):
    filename=safe_filename(filename);kind=EXTENSIONS.get(Path(filename).suffix.lower())
    if not kind: raise ExtractionError("Supported formats: PDF, DOCX, XLSX, CSV and TXT.")
    if not content: raise ExtractionError("The file is empty.")
    if len(content)>max_bytes: raise ExtractionError("The upload exceeds the configured size limit.")
    supplied=(mime or "").split(";")[0].lower()
    allowed={MIME[kind],"application/octet-stream",""}
    if kind=="CSV": allowed.update({"text/plain","application/csv","application/vnd.ms-excel"})
    if supplied not in allowed: raise ExtractionError("The supplied MIME type does not match the allowed format.")
    if kind=="PDF" and not content.startswith(b"%PDF-"): raise ExtractionError("The file is not a PDF.")
    if kind in {"DOCX","XLSX"}:
        with inspect_zip(content) as z:
            names=set(z.namelist())
            required="word/document.xml" if kind=="DOCX" else "xl/workbook.xml"
            if required not in names or "[Content_Types].xml" not in names: raise ExtractionError("The Office package does not match its extension.")
            if any("vbaProject" in name for name in names): raise ExtractionError("Macro-enabled content is not supported.")
    if kind in {"TXT","CSV"}:
        try: text=content.decode("utf-8-sig")
        except UnicodeDecodeError as exc: raise ExtractionError("Text/CSV uploads must use UTF-8 encoding.") from exc
        if "\x00" in text or text.lstrip().startswith(("%PDF-","PK\x03\x04")): raise ExtractionError("The file is not supported text content.")
    return filename,kind,MIME[kind]


def extract(content,kind,max_chunks=60,max_chars=120000):
    chunks=[];limitations=[];used=0
    def add(text,locator,chunk_type,**metadata):
        nonlocal used
        if not text.strip(): return
        if len(chunks)>=max_chunks or used>=max_chars:
            if "Extraction is partial: configured chunk/text limit reached." not in limitations: limitations.append("Extraction is partial: configured chunk/text limit reached.")
            return
        structured=metadata.get("structured_data")
        if structured is not None and len(json.dumps(structured,ensure_ascii=False))>6000:
            structured=None
            limitations.append("Large structured table/cell payload omitted; extracted text and original remain available.")
        # Preserve substrings; splitting does not invent additional pages/rows.
        for offset in range(0,len(text),6000):
            part=text[offset:offset+min(6000,max_chars-used)]
            if not part or len(chunks)>=max_chunks: 
                if "Extraction is partial: configured chunk/text limit reached." not in limitations: limitations.append("Extraction is partial: configured chunk/text limit reached.")
                break
            chunks.append(dict(chunk_type=chunk_type,source_locator=locator,text_content=part,sequence=len(chunks)+1,
                page_number=metadata.get("page_number"),sheet_name=metadata.get("sheet_name"),section_name=metadata.get("section_name"),structured_data=structured))
            used+=len(part)
    try:
        if kind=="PDF":
            pdf=PdfReader(BytesIO(content),strict=False)
            if pdf.is_encrypted: raise ExtractionError("Encrypted PDFs are not supported.")
            for index,page in enumerate(pdf.pages):
                if index>=200: limitations.append("PDF extraction stopped after 200 pages.");break
                try: text=page.extract_text(extraction_mode="layout") or ""
                except Exception: limitations.append(f"PDF page {index+1}: text extraction failed.");continue
                if text.strip(): add(text,f"PDF page {index+1}","pdf_page",page_number=index+1)
                else: limitations.append(f"PDF page {index+1}: no readable text; OCR is not available.")
            limitations.append("PDF tables retain extracted text layout; cell-level reconstruction is not guaranteed.")
        elif kind=="DOCX":
            with inspect_zip(content) as z:
                xml=z.read("word/document.xml")
                if b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper(): raise ExtractionError("DTD/entity declarations are not supported.")
                root=ET.fromstring(xml);ns={"w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                body=root.find("w:body",ns);paragraph=table=0;heading=None
                if body is None: raise ExtractionError("DOCX document body is missing.")
                text_of=lambda element: "".join(node.text or "" for node in element.findall(".//w:t",ns))
                for element in body:
                    if element.tag.endswith("}p"):
                        paragraph+=1;text=text_of(element)
                        style=element.find("w:pPr/w:pStyle",ns)
                        if style is not None and "heading" in style.get('{'+ns['w']+'}val','').lower(): heading=text[:255]
                        add(text,f"DOCX paragraph {paragraph}"+(f" · {heading}" if heading else ""),"docx_paragraph",section_name=heading)
                    elif element.tag.endswith("}tbl"):
                        table+=1
                        rows=[[text_of(cell) for cell in row.findall("w:tc",ns)] for row in element.findall("w:tr",ns)]
                        for start in range(0,len(rows),20):
                            block=rows[start:start+20]
                            add("\n".join(" | ".join(row) for row in block),f"DOCX table {table}, rows {start+1}–{start+len(block)}","docx_table",section_name=heading,structured_data={"rows":block})
                if any(name.startswith("word/embeddings/") for name in z.namelist()): limitations.append("Embedded DOCX objects are not extracted or executed.")
                limitations.append("DOCX headers, footers, tracked changes and rendered page numbers are not reconstructed.")
        elif kind=="XLSX":
            workbook=load_workbook(BytesIO(content),read_only=True,data_only=True,keep_links=False)
            for sheet in workbook:
                if sheet.max_row and sheet.max_row>10000: limitations.append(f"Sheet {sheet.title}: extraction limited to first 10000 rows.")
                if sheet.max_column and sheet.max_column>100: limitations.append(f"Sheet {sheet.title}: extraction limited to first 100 columns.")
                for row_index,row in enumerate(sheet.iter_rows(max_row=min(sheet.max_row or 0,10000),max_col=min(sheet.max_column or 0,100)),1):
                    populated=[cell for cell in row if cell.value is not None]
                    if not populated: continue
                    first,last=populated[0].column,populated[-1].column
                    values=[{"cell":cell.coordinate,"value":str(cell.value)} for cell in populated]
                    add("\n".join(f"{v['cell']}: {v['value']}" for v in values),f"XLSX sheet {sheet.title} · {get_column_letter(first)}{row_index}:{get_column_letter(last)}{row_index}","xlsx_row",sheet_name=sheet.title,structured_data={"cells":values})
                    if len(chunks)>=max_chunks: break
            workbook.close()
            limitations.append("Excel formulas are never evaluated; only stored/cached values are read. Uncached formulas and charts may be unavailable.")
        elif kind=="CSV":
            reader=csv.reader(StringIO(content.decode("utf-8-sig")),strict=True)
            try: headers=next(reader)
            except StopIteration: raise ExtractionError("CSV has no header row.")
            for row_index,row in enumerate(reader,2):
                if row_index>10000: limitations.append("CSV extraction limited to first 10000 records.");break
                if not row: continue
                add("Headers: "+" | ".join(headers)+"\nRow "+str(row_index)+": "+" | ".join(row),f"CSV row {row_index} (header row 1)","csv_row",structured_data={"headers":headers,"values":row,"record_row":row_index,"physical_end_line":reader.line_num})
                if len(row)!=len(headers): limitations.append(f"CSV row {row_index}: field count differs from header.")
                if len(chunks)>=max_chunks: break
        elif kind=="TXT":
            lines=content.decode("utf-8-sig").splitlines()
            for start in range(0,len(lines),20):
                add("\n".join(lines[start:start+20]),f"TXT lines {start+1}–{min(start+20,len(lines))}","text_lines")
        else: raise ExtractionError("Unsupported format.")
    except ExtractionError: raise
    except Exception as exc: raise ExtractionError(f"{kind} extraction failed; the original is retained.") from exc
    if len(chunks)>=max_chunks or used>=max_chars:
        limitations.append("Extraction is partial or at the configured chunk/text bound; review the original for additional material.")
    if not chunks: raise ExtractionError("No readable source text was extracted; the original is retained.",limitations)
    return chunks,list(dict.fromkeys(limitations))
